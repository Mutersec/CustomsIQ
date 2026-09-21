"""Import official EU Combined Nomenclature codes into the hs_codes table.

The CN reference file is published annually by the EU and must be downloaded
manually — this script never touches the network. Obtain it from the Eurostat
RAMON nomenclature server (https://ec.europa.eu/eurostat/ramon/) or the TARIC
consultation site (https://ec.europa.eu/taxation_customs/dds2/taric/), then:

    python scripts/import_cn_codes.py path/to/cn_codes.csv

Re-running the import is safe: hs_codes ends up with the latest value per
code either way, and a code whose description/category actually changed gets
a new entry in its version history (hs_code_history) rather than silently
losing the old value — see the "Importing the real CN nomenclature" section
of the README for how to inspect that history.
"""

import argparse
import csv
import logging
import re
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Optional

# Running this file directly puts scripts/ on sys.path, not the repo root, so
# the src.* imports below need the root prepended first.
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from src.customsiq.config import settings
from src.customsiq.database import (
    ImportStats,
    get_connection,
    record_cn_import,
    upsert_hs_codes_with_history,
)
from src.customsiq.logging_config import configure_logging
from src.customsiq.models import HSCode
from src.utils.validators import validate_cn_code

logger = logging.getLogger(__name__)

PROGRESS_EVERY = 1000
DEFAULT_BATCH_SIZE = 1000

# Header names seen across the RAMON and TARIC exports, lowercased.
CODE_COLUMNS = ("cn_code", "code", "goods code", "goods_code", "cn", "cncode")
DESCRIPTION_COLUMNS = (
    "description",
    "description_en",
    "description en",
    "self-explanatory text",
    "self_explanatory_text",
    "text",
)

# Hierarchy levels above an 8-digit CN leaf: chapter, heading, HS subheading.
# A real export is mostly these, so they are skipped quietly rather than warned about.
HIERARCHY_LENGTHS = (2, 4, 6)

# The 21 HS sections, which the CN inherits. Chapter 77 is reserved and 98-99
# are national/special use, so neither appears here — both fall back to "Other".
_SECTIONS: tuple[tuple[range, str], ...] = (
    (range(1, 6), "Animal Products"),
    (range(6, 15), "Vegetable Products"),
    (range(15, 16), "Fats & Oils"),
    (range(16, 25), "Food"),
    (range(25, 28), "Minerals"),
    (range(28, 39), "Chemicals"),
    (range(39, 41), "Plastics"),
    (range(41, 44), "Leather"),
    (range(44, 47), "Wood"),
    (range(47, 50), "Paper"),
    (range(50, 64), "Textile"),
    (range(64, 68), "Footwear"),
    (range(68, 71), "Stone & Glass"),
    (range(71, 72), "Precious Metals"),
    (range(72, 77), "Metals"),  # 77 is reserved by the HS, so it is left out
    (range(78, 84), "Metals"),
    (range(84, 86), "Machinery"),
    (range(86, 90), "Transport"),
    (range(90, 93), "Instruments"),
    (range(93, 94), "Arms"),
    (range(94, 97), "Miscellaneous"),
    (range(97, 98), "Art & Antiques"),
)

# Chapters specific enough to deserve their own label instead of their section's.
_CHAPTER_OVERRIDES = {
    "30": "Pharmaceutical",
    "85": "Electronics",
    "87": "Automotive",
    "94": "Furniture",
    "95": "Toys",
}

CHAPTER_CATEGORIES: dict[str, str] = {
    **{f"{chapter:02d}": category for chapters, category in _SECTIONS for chapter in chapters},
    **_CHAPTER_OVERRIDES,
}

UNKNOWN_CATEGORY = "Other"


class CNImportError(Exception):
    """Raised when the source file cannot be read as a CN export."""


def category_for_code(code: str) -> str:
    """Map a CN code to a category via its first two digits (the HS chapter).

    Args:
        code: A normalised, all-digit CN or TARIC code.

    Returns:
        The chapter's category, or "Other" for reserved/unknown chapters.
    """
    return CHAPTER_CATEGORIES.get(code[:2], UNKNOWN_CATEGORY)


def normalise_code(raw: str) -> str:
    """Strip the spaces, dots and dashes CN exports use to group digits."""
    return re.sub(r"[^0-9]", "", raw or "")


def resolve_column(headers: list[str], candidates: tuple[str, ...], override: Optional[str]) -> str:
    """Pick the column holding a given field, preferring an explicit override.

    Args:
        headers: The header row as read from the file.
        candidates: Known header aliases for this field, lowercased.
        override: A column name given on the command line, if any.

    Returns:
        The matching header, exactly as it appears in the file.

    Raises:
        CNImportError: If no candidate matches, listing the headers found.
    """
    if override:
        if override not in headers:
            raise CNImportError(f"Column {override!r} not found. File has: {', '.join(headers)}")
        return override
    for header in headers:
        if header.strip().lower() in candidates:
            return header
    raise CNImportError(
        f"Could not find a column among {', '.join(candidates)}. "
        f"File has: {', '.join(headers)}. Use --code-column / --description-column."
    )


def read_rows(path: Path) -> tuple[list[str], list[dict]]:
    """Read a CSV or Excel export into a header list and its rows.

    Rows are materialised rather than streamed: a full CN export is on the order
    of tens of thousands of rows, which costs little memory and lets the file be
    closed here instead of leaving a handle open for the caller.

    Excel support imports openpyxl lazily: it is not a project dependency, since
    the application never reads spreadsheets and only this tool would need it.

    Args:
        path: Path to a .csv/.tsv/.txt or .xlsx file.

    Returns:
        The header row, and the rows as {header: value} dicts.

    Raises:
        CNImportError: If the file has no header row, or Excel support is missing.
    """
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        return _read_excel_rows(path)

    with path.open(newline="", encoding="utf-8-sig") as handle:
        sample = handle.readline()
        handle.seek(0)
        try:
            dialect: type[csv.Dialect] = csv.Sniffer().sniff(sample, ",;\t")
        except csv.Error:
            dialect = csv.excel  # single-column or unsniffable file; comma is a safe default
        reader = csv.DictReader(handle, dialect=dialect)
        if not reader.fieldnames:
            raise CNImportError(f"{path} has no header row.")
        return list(reader.fieldnames), list(reader)


def _read_excel_rows(path: Path) -> tuple[list[str], list[dict]]:
    """Read an .xlsx export, if openpyxl happens to be installed."""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - depends on the environment
        raise CNImportError(
            "Excel support needs 'pip install openpyxl', "
            "or export the sheet to CSV and import that instead."
        ) from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = workbook.active.iter_rows(values_only=True)
        headers = [str(cell) if cell is not None else "" for cell in next(rows)]
        return headers, [
            {header: "" if value is None else str(value) for header, value in zip(headers, row)}
            for row in rows
        ]
    finally:
        workbook.close()


def parse_records(
    path: Path, code_column: Optional[str] = None, description_column: Optional[str] = None
) -> tuple[list[HSCode], int, int]:
    """Parse a CN export into HSCode records.

    Rows above the 8-digit leaf level (chapters, headings, subheadings) are
    skipped quietly — a real export is mostly those, so warning on them would
    bury the genuine problems. Rows that are neither a valid code nor a known
    hierarchy level are counted as malformed and logged.

    Args:
        path: The CN export to read.
        code_column: Explicit code column name, overriding auto-detection.
        description_column: Explicit description column name.

    Returns:
        The parsed records, the number of hierarchy rows skipped, and the
        number of malformed rows.
    """
    headers, rows = read_rows(path)
    code_key = resolve_column(headers, CODE_COLUMNS, code_column)
    description_key = resolve_column(headers, DESCRIPTION_COLUMNS, description_column)
    logger.info("reading %s (code=%r, description=%r)", path.name, code_key, description_key)

    records: list[HSCode] = []
    skipped = malformed = 0

    for line_number, row in enumerate(rows, start=2):
        code = normalise_code(row.get(code_key, ""))
        description = (row.get(description_key) or "").strip()

        if len(code) in HIERARCHY_LENGTHS:
            logger.debug("line %d: skipping hierarchy row %s", line_number, code)
            skipped += 1
            continue
        if not validate_cn_code(code) or not description:
            logger.warning(
                "line %d: skipping malformed row (code=%r, description=%r)",
                line_number,
                row.get(code_key, ""),
                description,
            )
            malformed += 1
            continue

        records.append(HSCode(code, description, category_for_code(code)))
        if len(records) % PROGRESS_EVERY == 0:
            logger.info("parsed %d codes…", len(records))

    return records, skipped, malformed


def import_file(
    path: Path,
    db_path: str,
    code_column: Optional[str] = None,
    description_column: Optional[str] = None,
    batch_size: int = DEFAULT_BATCH_SIZE,
    version_label: Optional[str] = None,
) -> int:
    """Parse a CN export and upsert it into the hs_codes table, versioning changes.

    Args:
        path: The CN export to import.
        db_path: SQLite database to write to.
        code_column: Explicit code column name, overriding auto-detection.
        description_column: Explicit description column name.
        batch_size: Rows written per upsert call.
        version_label: Label for this import run (e.g. "CN2026"), recorded
            against any code that changed. Defaults to a timestamp so every
            run is still identifiable if the caller doesn't name one.

    Returns:
        The number of codes imported.
    """
    records, skipped, malformed = parse_records(path, code_column, description_column)
    if version_label is None:
        version_label = f"import-{datetime.now(timezone.utc).isoformat()}"

    conn = get_connection(db_path)
    try:
        new_count = changed_count = unchanged_count = 0
        for start in range(0, len(records), batch_size):
            stats: ImportStats = upsert_hs_codes_with_history(
                conn, records[start : start + batch_size], version_label
            )
            new_count += stats.new_count
            changed_count += stats.changed_count
            unchanged_count += stats.unchanged_count
        record_cn_import(
            conn, version_label, path.name, datetime.now(timezone.utc).isoformat(), len(records)
        )
    finally:
        conn.close()

    logger.info(
        "imported %d codes into %s as %r (%d new, %d changed, %d unchanged, "
        "%d hierarchy rows skipped, %d malformed)",
        len(records),
        db_path,
        version_label,
        new_count,
        changed_count,
        unchanged_count,
        skipped,
        malformed,
    )
    return len(records)


def main(argv: Optional[list[str]] = None) -> int:
    """Command-line entry point.

    Args:
        argv: Argument list, defaulting to sys.argv[1:].

    Returns:
        0 on success, 1 if the file could not be read as a CN export.
    """
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("path", type=Path, help="Local CN reference file (.csv or .xlsx)")
    parser.add_argument("--db", default=settings.database_target, help="Target database")
    parser.add_argument("--code-column", help="Override the auto-detected code column")
    parser.add_argument("--description-column", help="Override the description column")
    parser.add_argument("--batch-size", type=int, default=DEFAULT_BATCH_SIZE)
    parser.add_argument(
        "--version-label", help="Label for this import run (e.g. CN2026); default is a timestamp"
    )
    args = parser.parse_args(argv)

    configure_logging()
    try:
        import_file(
            args.path,
            args.db,
            args.code_column,
            args.description_column,
            args.batch_size,
            args.version_label,
        )
    except (CNImportError, FileNotFoundError) as exc:
        logger.error("%s", exc)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
